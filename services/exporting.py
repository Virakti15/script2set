from io import BytesIO, StringIO
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from .schema import PRODUCTION, METADATA, normalize, ordered, text, UserError, production_view

def safe_cell(value):
    value='' if value is None else str(value)
    return "'"+value if value.lstrip().startswith(('=','+','-','@','\t','\r')) else value

def export_rows(rows):
    rows=ordered([normalize(r) for r in rows])
    if not rows: raise UserError('No shot breakdown available. Approve the shot plan and generate a breakdown first.')
    result=[]
    for r in rows:
        r['size_movement']=r['size_movement'] or ' / '.join(filter(None,[r['shot_size'],r['camera_movement']]))
        view=production_view(r)
        result.append([view[k] for k in PRODUCTION])
    return result

def csv_export(rows,metadata=None):
    buffer=StringIO(newline='')
    writer=csv.writer(buffer)
    writer.writerow(PRODUCTION.values())
    for r in export_rows(rows): writer.writerow([safe_cell(v) for v in r])
    return BytesIO(buffer.getvalue().encode('utf-8-sig'))

HEADER_ROW=6
DATA_ROW=7

def excel_export(rows,metadata=None,images=None,validation=None):
    wb=Workbook(); ws=wb.active; ws.title='Shot List'
    rows=ordered([normalize(r) for r in rows]);metadata=metadata or {}
    ws.append(['SHOT LIST']);ws.merge_cells('A1:S1');ws['A1'].font=Font(name='Arial',size=18,bold=True,color='172C29');ws.row_dimensions[1].height=32
    ws['A1'].fill=PatternFill('solid',fgColor='EAD1DC');ws['A1'].alignment=Alignment(horizontal='center',vertical='center')
    for row_index,keys in enumerate((('client','location','date'),('director','dop','duration'),('platform','aspect_ratio','reference')),2):
        for key,label_col,start,end in zip(keys,(1,5,9),(2,6,10),(4,8,19)):
            ws.cell(row_index,label_col,METADATA[key]);ws.cell(row_index,label_col).font=Font(name='Arial',size=11,bold=True)
            ws.cell(row_index,label_col).alignment=Alignment(wrap_text=True,vertical='center')
            cell=ws.cell(row_index,start,text(metadata.get(key)));cell.data_type='s';cell.font=Font(name='Arial',size=11)
            cell.alignment=Alignment(wrap_text=True,vertical='center');ws.merge_cells(start_row=row_index,start_column=start,end_row=row_index,end_column=end)
        ws.row_dimensions[row_index].height=30
        for cell in ws[row_index]:cell.fill=PatternFill('solid',fgColor='EAD1DC')
    ws['A5']='Setup colors follow scene order; thick borders start scenes. Shot comments retain source text. Amber assumptions require confirmation. VISUALS is optional.'
    ws.merge_cells('A5:S5');ws['A5'].font=Font(name='Arial',size=10,italic=True,color='555555');ws.row_dimensions[5].height=24
    header=HEADER_ROW
    for i,label in enumerate(PRODUCTION.values(),1):ws.cell(header,i,label)
    for r in export_rows(rows):
        ws.append(r)
        for cell in ws[ws.max_row]:cell.data_type='s'
        ws.cell(ws.max_row,1).number_format='@'
    ws.freeze_panes=f'D{DATA_ROW}';ws.auto_filter.ref=f'A{HEADER_ROW}:S{ws.max_row}'
    for cell in ws[header]:
        cell.font=Font(name='Arial',color='172C29',bold=True,size=11); cell.fill=PatternFill('solid',fgColor='D9EAD3')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[header].height=42
    scene_index=-1;last_scene=None;last_setup=None
    widths=[13,15,28,23,20,25,20,13,48,28,52,38,14,22,25,30,30,40,45]
    palette=['D9EAD3','DCEAF7','FFF2CC','F4CCCC','EAD1DC']
    for row,source in zip(ws.iter_rows(min_row=DATA_ROW),rows):
        new_scene=source['scene']!=last_scene
        setup=tuple(source[k] for k in ('scene','shot_location','size_angle','camera_movement','lighting'))
        if setup!=last_setup:scene_index+=1;last_setup=setup
        last_scene=source['scene']
        for cell in row:
            cell.font=Font(name='Arial',size=11,color='24332E'); cell.alignment=Alignment(wrap_text=True,vertical='center')
            cell.fill=PatternFill('solid',fgColor=palette[scene_index%len(palette)])
            cell.border=Border(bottom=Side(style='hair',color='CAD4CF'),top=Side(style='medium' if new_scene else 'hair',color='78948A' if new_scene else 'CAD4CF'))
        lines=max(sum(max(1,(len(line)+width-2)//(width-1)) for line in str(cell.value or '').split('\n')) for cell,width in zip(row,widths))
        ws.row_dimensions[row[0].row].height=min(409,max(96,lines*16+16))
        row[0].comment=Comment(f'Scene {source["scene"]}: {source.get("scene_name","")}\nBeat {source.get("beat","")}\n'+source.get('source_provenance','')+'\n'+source.get('source_script_block',''),'Script2Set')
        warnings=[v['message'] for v in (validation or {}).get('warnings',[]) if v.get('row_id')==source['id'] or not v.get('row_id')]
        if warnings:
            row[18].comment=Comment('\n\n'.join(warnings),'Script2Set')
            row[18].fill=PatternFill('solid',fgColor='FFE2A6')
        image_data=(images or {}).get(source['id'])
        if image_data:
            picture=ExcelImage(BytesIO(image_data));scale=min(180/picture.width,110/picture.height)
            picture.width=int(picture.width*scale);picture.height=int(picture.height*scale)
            ws.add_image(picture,f'C{row[0].row}')
            ws.row_dimensions[row[0].row].height=max(96,ws.row_dimensions[row[0].row].height)
    for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.print_title_rows=f'{HEADER_ROW}:{HEADER_ROW}'
    ws.print_area=f'A1:S{ws.max_row}'
    ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=80
    ws.print_options.horizontalCentered=True
    out=BytesIO(); wb.save(out); out.seek(0); return out

def ppt_export(rows, per_slide=1, title='Script2Set', images=None, metadata=None):
    if per_slide not in (1,2): raise UserError('Choose one or two shots per slide.')
    rows=ordered([normalize(r) for r in rows])
    if not rows: raise UserError('There are no shots to export.')
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    def textbox(slide,x,y,w,h,value,size=14,color='DDE8E4',bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=box.text_frame; tf.word_wrap=True
        tf.margin_left=0; tf.margin_top=0
        # Bound slide text visually; preserve the complete text in speaker notes.
        p=tf.paragraphs[0]; p.text=str(value)
        p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        return box
    for start in range(0,len(rows),per_slide):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor.from_string('14231F')
        textbox(slide,.5,.25,11,.3,title[:80].upper(),12,'9EBAB0',True)
        textbox(slide,12,.25,1,.3,f'{start//per_slide+1:02}',12,'9EBAB0')
        notes=[]
        for j,r in enumerate(rows[start:start+per_slide]):
            y=.95+j*3.2 if per_slide==2 else 1.05
            h=2.7 if per_slide==2 else 4.8
            textbox(slide,.5,y,6,.4,f'SHOT {r["number"]} · SCENE {r["scene"]}',15,'FFFFFF',True)
            frame_y=y+.55; frame_h=h-.55
            image_data=(images or {}).get(r['id'])
            if image_data:
                pic=slide.shapes.add_picture(BytesIO(image_data),Inches(.5),Inches(frame_y),width=Inches(6.3))
                # Contain with preserved aspect ratio.
                if pic.height>Inches(frame_h):
                    ratio=Inches(frame_h)/pic.height; pic.width=int(pic.width*ratio); pic.height=Inches(frame_h)
                pic.left=Inches(.5)+(Inches(6.3)-pic.width)//2
            else:
                shape=slide.shapes.add_shape(1,Inches(.5),Inches(frame_y),Inches(6.3),Inches(frame_h))
                shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string('203B32')
                shape.line.color.rgb=RGBColor.from_string('4D7566')
                textbox(slide,.85,frame_y+.25,5.6,.4,r['shot_size'] or 'Frame to confirm',18,'C4DCCD',True)
                textbox(slide,.85,frame_y+.8,5.6,.75,r['shot_description'][:180],13)
                textbox(slide,.85,frame_y+frame_h-.4,5.6,.3,'STORYBOARD IMAGE PENDING',10,'9EBAB0')
            x=7.2; max_script=280 if per_slide==2 else 750
            textbox(slide,x,y,5.5,.4,r['shot_type'] or 'LIVE ACTION',12,'97CFAE',True)
            textbox(slide,x,y+.45,5.5,.7 if per_slide==2 else 1.2,r['shot_description'][:180 if per_slide==2 else 360],15,'FFFFFF')
            textbox(slide,x,y+(1.2 if per_slide==2 else 1.8),5.5,.8 if per_slide==2 else 2,
                    r['script'][:max_script]+('… (full text in notes)' if len(r['script'])>max_script else ''),11 if per_slide==2 else 13)
            textbox(slide,x,y+h-.4,5.5,.6,f'{r["shot_size"]} · {r["shot_angle"]} · {r["camera_movement"]}'[:150],10,'9EBAB0')
            notes.append('\n'.join(f'{k}: {v}' for k,v in r.items() if k not in ('existing','image')))
        slide.notes_slide.notes_text_frame.text=('\n'.join(f'{METADATA.get(k,k)}: {v}' for k,v in (metadata or {}).items())+'\n\n' if start==0 else '')+'\n\n'.join(notes)
        textbox(slide,.5,7.15,12,.2,'SCRIPT2SET   /   FROM SCRIPT TO SHOOT-READY',9,'7F9C90')
    out=BytesIO(); prs.save(out); out.seek(0); return out
