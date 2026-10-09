from io import BytesIO
from zipfile import ZipFile, BadZipFile
import re
from openpyxl import load_workbook
from .schema import PRODUCTION, INTERNAL, METADATA, PLAN_FIELDS, ALIASES, normalize, UserError, text

def key(s):
    return re.sub(r'[^a-z0-9]','',str(s).lower())

MAPPING={key(v):k for k,v in {**INTERNAL,**PRODUCTION}.items()}
MAPPING.update({key(k):k for k in [*PLAN_FIELDS,*PRODUCTION,'scene','number',*ALIASES]})
MAPPING.update({key(k):v for k,v in ALIASES.items()})
MAPPING.update({'shotno':'number','shot':'number','shotnumber':'number','sceneno':'scene','scenenumber':'scene',
                'daynightintext':'day_int_ext',
                'dialogue':'script','description':'shot_description','location':'shot_location','movement':'camera_movement',
                'angle':'shot_angle','size':'shot_size','wardrobe':'costumes','characters':'cast'})
MAPPING.update({'shotsizecameramovement':'size_movement','shotlocationrequirementsetupwise':'location_setup',
                'designelements':'design_elements','artsetup':'art_setup','costumes':'costumes','script':'script',
                'beat':'beat','scenesequencename':'scene_name'})

def import_excel(data):
    if not data or len(data)>3_000_000:
        raise UserError('Choose an .xlsx file smaller than 3 MB.')
    try:
        with ZipFile(BytesIO(data)) as z:
            if sum(f.file_size for f in z.infolist())>30_000_000 or len(z.infolist())>2000:
                raise UserError('This workbook is too large when expanded. Export only the relevant shot sheet.')
        wb=load_workbook(BytesIO(data),read_only=True,data_only=False)
    except UserError: raise
    except (BadZipFile,ValueError,KeyError,OSError):
        raise UserError('This is not a readable Excel workbook. Upload a valid .xlsx file.') from None
    try:
        best=None
        for sheet in wb.worksheets:
            for idx,values in enumerate(sheet.iter_rows(max_row=20,max_col=60,values_only=True),1):
                mapping={i:MAPPING[key(v)] for i,v in enumerate(values) if v is not None and key(v) in MAPPING}
                score=len(set(mapping.values()))
                if ('script' in mapping.values() or 'shot_description' in mapping.values()) and (best is None or score>best[0]):
                    best=(score,sheet,idx,mapping)
        if best is None:
            raise UserError('No shot worksheet was found. Include a Script, Dialogue, or Shot Description column.')
        _,sheet,header,mapping=best
        rows=[]; warnings=[];metadata={};scene_number=0;last_setup=None
        for values in sheet.iter_rows(max_row=header-1,values_only=True):
            for i,value in enumerate(values[:-1]):
                field=next((k for k,v in METADATA.items() if key(value) in (key(v),key(k))),None)
                if field and values[i+1] is not None:metadata[field]=text(values[i+1])
        for values in sheet.iter_rows(min_row=header+1,max_row=header+1001,max_col=60,values_only=True):
            if not any(v is not None for v in values): continue
            row={}
            for col,field in mapping.items():
                value=values[col]
                if isinstance(value,str) and value.startswith('='):
                    warnings.append('Formula cells were skipped; paste their values before importing to retain them.'); continue
                if value is not None:
                    if field=='number' and hasattr(value,'year'):raise UserError('A SHOT cell contains an Excel date. Restore the original shot identifier as text before importing.')
                    row[field]=str(value) if field=='source_script_block' else text(value)
            if not row.get('script') and not row.get('shot_description'): continue
            if len(rows)>=300:
                raise UserError('This sheet has more than 300 shots. Split it into smaller projects.')
            setup=(row.get('shot_location',''),row.get('day_int_ext',''))
            if setup!=last_setup:scene_number+=1;last_setup=setup
            row['scene']=row.get('scene') or (row['number'].split('.')[0] if '.' in row.get('number','') else str(scene_number or 1))
            row['number']=row.get('number') or f'{row["scene"]}.{1+sum(r["scene"]==row["scene"] for r in rows)}'
            row['existing']={k:v for k,v in row.items() if k in (set(PRODUCTION)|set(INTERNAL)) and v}
            row['production_order']=len(rows)
            if row.get('size_angle'):
                parts=re.split(r'\s*/\s*',row['size_angle'],maxsplit=1);row['shot_size']=parts[0];row['shot_angle']=parts[1] if len(parts)>1 else ''
                row['size_angle']=' / '.join(parts)
                row['existing']['size_angle']=row['size_angle']
            if row.get('day_int_ext'):
                parts=re.split(r'\s*/\s*',row['day_int_ext'],maxsplit=1);row['day_night']=parts[0];row['interior_exterior']=parts[1] if len(parts)>1 else ''
                row['day_int_ext']=' / '.join(parts)
                row['existing']['day_int_ext']=row['day_int_ext']
            row['source_script_block']=row.get('source_script_block') or row.get('script') or row.get('shot_description','')
            row['source_provenance']='Imported worksheet source text; exact cell value'
            row['shot_description']=row.get('shot_description') or row.get('script','')
            row['shot_type']=row.get('shot_type') or ('Title Card / Graphic' if 'graphic' in row.get('shot_size','').lower() else 'Live Action')
            for field in ('shot_size','shot_angle','camera_movement','cast','props','shot_location','day_night','interior_exterior'):
                row.setdefault(field,'Confirmation Required')
            if row.get('size_movement'):
                parts=row['size_movement'].split(' / ',1)
                row['shot_size']=parts[0]; row['camera_movement']=parts[1] if len(parts)>1 else 'Confirmation Required'
            rows.append(normalize(row))
        if not rows: raise UserError('The selected worksheet contains no usable shot rows.')
        return {'shots':rows,'worksheet':sheet.title,'metadata':metadata,'warnings':list(dict.fromkeys(warnings))}
    finally:
        wb.close()
