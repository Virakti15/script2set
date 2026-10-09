from copy import deepcopy
import pytest
from services.accuracy import audit
from test_architecture import CORPUS, broken_down
from test_workflow import client

@pytest.fixture(autouse=True)
def offline(monkeypatch):monkeypatch.setenv('AI_PROVIDER','demo')

@pytest.mark.parametrize('style',CORPUS)
def test_original_script_coverage(style):
    script=CORPUS[style];_,plan,rows=broken_down(script)
    report=audit(rows,plan,script)
    assert report['ok'],report['errors']
    assert report['source_verified']

def test_detects_missing_and_changed_dialogue():
    script=CORPUS['dialogue'];_,plan,rows=broken_down(script)
    rows[0]['script']='Aarti: Invented words.'
    codes={v['code'] for v in audit(rows,plan,script)['errors']}
    assert {'dialogue_missing','dialogue_invented'}<=codes
    assert 'uncovered_beat' in {v['code'] for v in audit([],plan,script)['errors']}

def test_source_and_numbering_integrity():
    script=CORPUS['locations'];_,plan,rows=broken_down(script)
    rows[1]['number']=rows[0]['number'];rows[0]['source_script_block']='Changed source'
    codes={v['code'] for v in audit(rows,plan,script)['errors']}
    assert {'shot_overlap','source_changed'}<=codes

def test_optional_visuals_and_unconfirmed_choices():
    script=CORPUS['dialogue'];_,plan,rows=broken_down(script)
    assert not rows[0]['visuals']
    rows[0]['equipment']='Expensive crane';rows[0]['shot_location']='Invented castle'
    report=audit(rows,plan,script)
    assert report['ok']
    assert sum(v['code']=='unsupported_detail' for v in report['warnings'])>=2

def test_explicit_framing_conflict():
    _,plan,rows=broken_down(CORPUS['dialogue'])
    rows[0].update(shot_size='ECU',size_angle='ECU / Top',shot_description='Wide of the office')
    assert 'framing_conflict' in {v['code'] for v in audit(rows,plan)['errors']}

def test_silent_cast_props_and_wardrobe_boundaries():
    _,_,rows=broken_down('INT. KITCHEN - DAY\nMS of Mira placing a saucepan on the stove.\nMira wears a navy jacket and sets a red folder on the desk.\nCU of Mira lifting the cup.')
    assert all(r['cast']=='Mira' for r in rows)
    assert any('folder' in r['props'] for r in rows)
    assert 'sets' not in rows[-1]['costumes'] and 'navy jacket' in rows[-1]['costumes']

def test_wide_direction_preserved_with_speaking_character():
    _,_,rows=broken_down('EXT. STATION - DAY\nMira: Hello.\nWIDE: Mira walks to the platform. Handheld camera.')
    assert rows[-1]['shot_size']=='Wide' and rows[-1]['camera_movement']=='Handheld'

def test_excel_route_blocks_uncovered_source(client):
    from test_workflow import post
    script=CORPUS['locations'];_,plan,rows=broken_down(script)
    plan=plan[:1];rows=rows[:1]
    approval=post(client,'/api/approve',{'plan':plan,'reviewed':True}).json
    data={'plan':plan,'breakdown':rows,'approval':approval['approval'],'breakdown_hash':approval['plan_hash'],'source_script':script}
    result=post(client,'/api/export/xlsx',data)
    assert result.status_code==422
    assert 'uncovered_beat' in {e['code'] for e in result.json['validation']['errors']}

def test_reference_compact_headers_and_slashes_survive_regeneration():
    from io import BytesIO
    from openpyxl import Workbook
    from services.importing import import_excel
    from services.planning import generate_breakdown
    wb=Workbook();ws=wb.active
    ws.append(['Client:','Example',None,None,'Location:','Studio',None,None,'Date:','TBD'])
    ws.append(['SHOT','DAY/NIGHT\nINT/EXT','SHOT SIZE/\nSHOT ANGLE','SHOT DESCRIPTION','SCRIPT / DIALOGUE'])
    ws.append(['6–12 Process','DAY/INT','ECU/Top','ECU of a phone','Phone process'])
    data=BytesIO();wb.save(data);imported=import_excel(data.getvalue())
    assert imported['metadata']=={'client':'Example','location':'Studio','date':'TBD'}
    row=generate_breakdown(imported['shots'],'1')[0]
    assert row['shot_size']=='ECU' and row['shot_angle']=='Top'
    assert row['day_night']=='DAY' and row['interior_exterior']=='INT'

def test_setup_bands_and_scene_borders():
    from openpyxl import load_workbook
    from services.exporting import excel_export
    _,_,rows=broken_down('INT. ROOM - DAY\nMira: Hello.\nREACTION: Mira smiles.\nEXT. PARK - DAY\nMira walks.')
    ws=load_workbook(excel_export(rows)).active
    assert ws['A7'].fill.fgColor.rgb!=ws['A8'].fill.fgColor.rgb
    assert ws['A7'].border.top.style==ws['A9'].border.top.style=='medium'
    assert ws['A8'].border.top.style=='hair'

@pytest.mark.parametrize('description',['"','Mira:','Confirmation Required'])
def test_meaningless_descriptions_block_export(description):
    _,plan,rows=broken_down(CORPUS['dialogue']);rows[0]['shot_description']=description
    assert 'empty_setup' in {e['code'] for e in audit(rows,plan)['errors']}

def test_screen_documents_are_not_physical_props():
    from services.beats import physical_props
    assert physical_props('Laptop workflow in the same screen setup. Open the brief.')==['Laptop']
    assert 'brief' in physical_props('In the same screen setup, place a printed brief beside the laptop.')

def test_prop_names_are_unique_ignoring_capitalization():
    from services.beats import physical_props
    assert physical_props('Phone POV: the phone interface.')==['Phone']
