"""Regression corpus: deliberately different screenplay styles, never one example."""
from copy import deepcopy
from io import BytesIO, StringIO
import csv
import json
import pytest
from openpyxl import load_workbook, Workbook
from pptx import Presentation
from services.schema import PRODUCTION, normalize, ordered, fingerprint, production_view, validate_plan, UserError
from services.screenplay import analyze
from services.planning import generate_plan, generate_breakdown
from services.importing import import_excel
from services.exporting import excel_export, csv_export, ppt_export
from test_workflow import client, post

CORPUS={
 'dialogue': 'INT. OFFICE - DAY\nAarti: What happened?\nManik: Nothing.\nAarti: Are you sure?\nManik: Yes.',
 'action': 'INT. GARAGE - NIGHT\nRavi enters and puts his bag down.\nRavi turns toward the door.\nRavi slams the door.\nREACTION: Ravi freezes at the sound.',
 'locations': 'INT. KITCHEN - DAY\nMira opens a cupboard.\n\nEXT. GARDEN - DAY\nMira walks to the gate.\n\nEXT. GARDEN - NIGHT\nMira holds a lamp.',
 'flashback': 'INT. STUDIO - DAY\nMira: I remember it.\n\n2 minutes earlier\nMira drops the letter.\n\nBACK TO PRESENT\nMira: Now you know.',
 'cards': 'TITLE CARD: OUR STORY\nSTAGE 2: ANGER\nON-SCREEN TEXT: Tomorrow\nSTAGE 6: DELEGATION',
 'montage': 'MONTAGE: Morning routine\n- Mira opens the fridge.\n- CLOSE-UP: a phone alarm on a desk.\n- Ravi runs on a bridge.\nEND MONTAGE',
 'vo': 'EXT. BEACH - DAWN\nWaves reach the sand.\nNARRATOR (V.O.): Every morning is a new beginning.\nMUSIC: quiet strings',
 'pov_split': 'INT. OFFICE - DAY\nPOV: Aarti sees the red phone.\nSPLIT SCREEN:\nLEFT: INT. OFFICE - DAY\nAarti: Hello.\nRIGHT: EXT. PARK - DAY\nManik: Can you hear me?',
 'sound': 'INT. HALL - NIGHT\nMira opens a door.\nSFX: door creaks\nMUSIC: ominous drone\nMira: Who is there?',
 'messy': '\"\r\n  \"Mira:  \r\n\r\n Hello there\"\r\n\r\n   Ravi : Fine.\r\n\r\nSTAGE 1: BEGIN\r\n',
 'long': 'INT. OFFICE - DAY\n'+ '\n'.join(('Mira: '+('We need to review the plan carefully. '*20),'Ravi: '+('Every detail matters to the team. '*20))*8),
 'props': 'INT. CONFERENCE ROOM - DAY\nAarti wears a navy jacket.\nLIGHTING: soft window light from camera left.\nAarti: Are we ready?\nAarti slams the campaign brief onto the conference table.\nINSERT: the red phone lights up.\nREACTION: Aarti stares at the phone.',
 'speed_vfx': 'EXT. ROOFTOP - NIGHT\nMira runs to the ledge.\nSLOW MOTION: a glass bottle falls.\nNORMAL SPEED: Mira turns.\nVFX: a digital city dissolves behind her.',
 'compilation': 'COMPILATION: A working day\n- A baker opens the oven.\n- A cyclist crosses a bridge.\n- A customer holds a cup.\nEND COMPILATION',
}

@pytest.fixture(autouse=True)
def demo(monkeypatch):monkeypatch.setenv('AI_PROVIDER','demo')

def planned(script):
    analysis=analyze(script);rows=[]
    for chunk in analysis['chunks']:rows+=generate_plan(chunk,analysis['characters'],offset=len(rows),previous=rows[-4:])
    return analysis,rows

def broken_down(script):
    analysis,plan=planned(script)
    rows=[]
    for s in dict.fromkeys(r['scene'] for r in plan):rows+=generate_breakdown(plan,s)
    return analysis,plan,rows

@pytest.mark.parametrize('style',CORPUS)
def test_corpus_hierarchy_source_order_and_exports(style):
    script=CORPUS[style];analysis,plan,rows=broken_down(script)
    assert plan and len(plan)==len(rows)
    validate_plan(plan)
    assert len({r['number'] for r in plan})==len(plan)
    assert [r['number'] for r in ordered(plan)]==[r['number'] for r in plan]
    assert [r['source_order'] for r in plan]==sorted(r['source_order'] for r in plan)
    for r in rows:
        assert r['number'] and r['beat'].startswith(r['scene']+'.')
        assert r['source_script_block']==''.join(script[s['start']:s['end']] for s in r['source_spans'])
        assert r['source_script_block'].strip() not in ('"','Mira:','Ravi:')
        assert all(k in r for k in PRODUCTION) and r['clip_no']==''
    wb=load_workbook(excel_export(rows));sheet=wb.active
    assert sheet.max_column==19 and sheet.max_row==len(rows)+6 and all(r.max_row<6 for r in sheet.merged_cells.ranges)
    assert list(list(sheet.values)[5])==list(PRODUCTION.values())
    sheet_rows=list(sheet.values)[6:]
    for i,(source,exported) in enumerate(zip(ordered(rows),sheet_rows)):
        assert exported[0]==source['number']
        assert list(exported)==[production_view(source)[k] or None for k in PRODUCTION]
    csv_rows=list(csv.DictReader(StringIO(csv_export(rows).getvalue().decode('utf-8-sig'))))
    assert len(csv_rows)==len(rows) and list(csv_rows[0])==list(PRODUCTION.values())
    deck=Presentation(ppt_export(rows,2))
    assert len(deck.slides)==(len(rows)+1)//2

def test_exchange_is_one_shootable_master():
    _,rows=planned(CORPUS['dialogue'])
    assert len(rows)==1 and rows[0]['shot_size']=='Two Shot'
    assert rows[0]['cast']=='Aarti, Manik'

def test_small_actions_group_but_slam_and_reaction_split():
    _,rows=planned(CORPUS['action'])
    assert len(rows)==3
    assert 'puts his bag' in rows[0]['source_script_block'] and 'turns toward' in rows[0]['source_script_block']
    assert rows[-1]['shot_type']=='Reaction'

def test_scene_boundaries_are_contextual():
    analysis,rows=planned(CORPUS['locations'])
    assert len(analysis['scenes'])==3
    assert [r['shot_location'] for r in rows]==['KITCHEN','GARDEN','GARDEN']
    assert [r['day_night'] for r in rows]==['DAY','DAY','NIGHT']

def test_blank_lines_and_repeated_same_heading_do_not_make_scenes():
    data=analyze('INT. OFFICE - DAY\nMira: Hello\n\n\nRavi: Hi\n\nINT. OFFICE - DAY\nMira: Again')
    assert len(data['scenes'])==1

def test_flashback_and_return_restore_context():
    data,rows=planned(CORPUS['flashback'])
    assert len(data['scenes'])==3
    assert [r['temporal_context'] for r in rows]==['Present','2 minutes earlier','Present']
    assert all(r['shot_location']=='STUDIO' for r in rows)

def test_flashback_heading_does_not_create_empty_scene():
    data,rows=planned('INT. OFFICE - DAY\nMira: Hello\nFLASHBACK\nEXT. GARDEN - NIGHT\nMira runs.\nBACK TO PRESENT\nMira: Fine')
    assert len(data['scenes'])==3 and [r['scene'] for r in rows]==['1','2','3']
    assert rows[-1]['shot_location']=='OFFICE' and rows[-1]['day_night']=='DAY'

def test_cards_are_production_graphics():
    _,_,rows=broken_down(CORPUS['cards'])
    assert len(rows)==4
    for r in rows:
        assert r['cast']=='None' and r['props']=='None' and r['camera_movement']=='Static'
        assert r['shot_size']=='Full Frame Graphic'
    assert rows[-1]['design_elements']=='STAGE 6: DELEGATION'

@pytest.mark.parametrize('style',['montage','compilation'])
def test_montage_has_separate_visual_setups(style):
    _,rows=planned(CORPUS[style])
    assert len(rows)==3
    assert len({r['beat'] for r in rows})==3

def test_vo_is_heard_not_visible():
    _,_,rows=broken_down(CORPUS['vo'])
    assert all('Narrator' not in r['cast'].title() for r in rows)
    assert 'NARRATOR' in rows[0]['sound'] and 'quiet strings' in rows[0]['sound']

def test_split_screen_is_two_plates_under_one_beat():
    _,_,rows=broken_down(CORPUS['pov_split'])
    assert rows[0]['shot_angle']=='POV'
    panels=[r for r in rows if r['shot_type']=='Split Screen']
    assert len(panels)==2 and panels[0]['beat']==panels[1]['beat']
    assert [r['shot_location'] for r in panels]==['OFFICE','PARK']
    assert [r['cast'] for r in panels]==['Aarti','Manik']
    assert all('compositing' in r['vfx'] for r in panels)

def test_sfx_and_music_do_not_become_dialogue_rows():
    _,_,rows=broken_down(CORPUS['sound'])
    assert len(rows)==1
    assert rows[0]['sfx']=='SFX: door creaks'
    assert rows[0]['sound']=='MUSIC: ominous drone'
    assert 'SFX:' not in rows[0]['script']

def test_very_long_single_dialogue_has_bounded_source_continuations():
    source='Mira: '+('This sentence continues the same thought. '*200)
    _,rows=planned(source)
    assert 1<len(rows)<10
    assert all(len(r['source_script_block'])<=2400 for r in rows)
    assert ''.join(r['source_script_block'] for r in rows)==source

def test_prop_action_context_is_coherent():
    _,_,rows=broken_down(CORPUS['props'])
    slam=next(r for r in rows if 'slams' in r['shot_description'])
    assert slam['cast']=='Aarti' and 'campaign brief' in slam['props']
    assert slam['shot_location']=='CONFERENCE ROOM' and 'navy jacket' in slam['costumes']
    assert 'soft window light' in slam['lighting'] and 'campaign brief' in slam['continuity']

def test_montage_named_setup_location_overrides_previous_scene():
    script='INT. CONFERENCE ROOM - DAY\nAarti: Hello\nEXT. COURTYARD - NIGHT\nMira: Wait\nMONTAGE: Search\n- Aarti opens a cupboard in the conference room.\n- Mira runs across the courtyard.\n- INSERT: a phone on a desk.'
    _,rows=planned(script);montage=[r for r in rows if r['shot_type']=='Montage setup']
    assert [r['shot_location'] for r in montage]==['CONFERENCE ROOM','COURTYARD','Confirmation Required']

def test_pov_owner_is_required_for_context_but_not_visible():
    _,rows=planned('INT. ROOM - DAY\nMira: Hello\nPOV: Mira sees a letter.')
    assert rows[-1]['cast']=='None'
    assert 'Mira' in rows[-1]['scene_context']['cast_present']

def test_wardrobe_and_lighting_changes_only_flow_forward():
    script='INT. ROOM - DAY\nMira wears a red jacket.\nLIGHTING: warm window light.\nMira: Hello.\nINSERT: a letter.\nMira changes into a blue jacket.\nLIGHTING: changes to cool overhead light.\nMira: Ready.'
    _,_,rows=broken_down(script)
    assert 'red jacket' in rows[0]['costumes'] and 'blue jacket' not in rows[0]['costumes']
    assert 'warm window light' in rows[0]['lighting'] and 'cool overhead' not in rows[0]['lighting']
    assert 'blue jacket' in rows[-1]['costumes'] and 'cool overhead' in rows[-1]['lighting']

def test_speed_and_vfx_keep_separate_execution():
    _,_,rows=broken_down(CORPUS['speed_vfx'])
    assert len(rows)==4
    assert 'Slow motion' in rows[1]['fps'] and rows[2]['fps']=='24 FPS'
    assert 'digital city' in rows[-1]['vfx']

def test_source_is_immutable_through_breakdown_and_exports(client):
    script='  "Mira:  \r\n Hello"\r\n\r\nTITLE CARD: END'
    _,plan=planned(script);approval=post(client,'/api/approve',{'plan':plan,'reviewed':True}).json
    rows=generate_breakdown(plan,'1');rows[0]['script']='Manual dialogue edit';rows[0]['source_script_block']='Tampered source'
    data={'plan':plan,'approval':approval['approval'],'breakdown_hash':approval['plan_hash'],'breakdown':rows}
    response=post(client,'/api/export/xlsx',data);assert response.status_code==200
    sheet=load_workbook(BytesIO(response.data)).active
    assert sheet['K7'].value=='Manual dialogue edit'
    printed=post(client,'/api/print',data).json
    assert printed['shots'][0]['source_script_block']==plan[0]['source_script_block']

def test_migration_retains_legacy_manual_details(client):
    old={'id':'fixed','scene':'1','number':'1.10','script':'Mira: Hello','visual_action':'Mira leans forward','suggested_shot_size':'Wide','size_movement':'Wide / Pan'}
    result=post(client,'/api/migrate',{'project':{'version':1,'plan':[old],'breakdown':[{**old,'costumes':'Green jacket','lighting':'Softbox left'}]}})
    p=result.json['project'];assert p['version']==3 and p['stale']
    assert p['plan'][0]['number']=='1.10' and p['plan'][0]['beat']=='1.1'
    assert p['plan'][0]['shot_size']=='Wide' and p['plan'][0]['camera_movement']=='Pan'
    assert generate_breakdown(p['plan'],'1')[0]['costumes']=='Green jacket'

def test_numeric_hierarchy_across_scenes_beats_and_shots():
    rows=[normalize({'scene':s,'beat':b,'number':n,'script':'x'}) for s,b,n in [('2','2.1','2.1.1'),('1','1.10','1.10.1'),('1','1.2','1.2.11'),('1','1.2','1.2.2'),('1','1.2','1.2.10')]]
    assert [r['number'] for r in ordered(rows)]==['1.2.2','1.2.10','1.2.11','1.10.1','2.1.1']

def test_export_import_roundtrip_all_19_columns():
    _,_,rows=broken_down(CORPUS['props'])
    result=import_excel(excel_export(rows).getvalue())
    assert len(result['shots'])==len(rows)
    for source,imported in zip(rows,result['shots']):
        for k in PRODUCTION:assert imported[k]==source[k]

def test_cloud_can_return_multiple_shots_in_one_beat(monkeypatch):
    scene=analyze(CORPUS['dialogue'])['chunks'][0];beat=scene['beats'][0]
    monkeypatch.setenv('AI_PROVIDER','ollama')
    proposed=[{'beat':beat['number'],'source_ids':beat['source_ids'],'shot_description':'Two shot holding the exchange','shot_size':'Two Shot','shot_purpose':'Establish shared space'},
              {'beat':beat['number'],'source_ids':[beat['source_ids'][-1]],'shot_description':'Close-up of Manik answering','shot_size':'Close-Up','shot_purpose':'Isolate the final reassurance'}]
    monkeypatch.setattr('services.providers.post_json',lambda *a,**k:{'message':{'content':json.dumps({'shots':proposed})}})
    rows=generate_plan(scene)
    assert [r['number'] for r in rows]==['1','2'] and rows[0]['beat']==rows[1]['beat']

def test_provider_cannot_rewrite_source_or_numbering(monkeypatch):
    scene=analyze('Mira: Hello')['chunks'][0];beat=scene['beats'][0]
    monkeypatch.setenv('AI_PROVIDER','ollama')
    row={'beat':beat['number'],'source_ids':beat['source_ids'],'source_script_block':'Invented','number':'9.9.9','shot_description':'Mira waits','shot_size':'Medium'}
    monkeypatch.setattr('services.providers.post_json',lambda *a,**k:{'message':{'content':json.dumps({'shots':[row]})}})
    result=generate_plan(scene)[0]
    assert result['source_script_block']=='Mira: Hello' and result['number']=='1'
