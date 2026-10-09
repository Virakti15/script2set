"""New providers/storage/storyboard/hosting tests. Network calls are explicitly mocked."""
import base64
import copy
import json
import os
import shutil
import subprocess
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
import pytest
from PIL import Image
from pptx import Presentation
from app import create_app
from services import providers, storyboard
from services.schema import UserError
from test_workflow import post, fixture_plan

@pytest.fixture
def client(monkeypatch,tmp_path):
    for key in ('APP_ENV','VERCEL','APP_ACCESS_PASSWORD','FRONTEND_ORIGINS','OPENAI_API_KEY','IMAGE_API_KEY','IMAGE_MODEL'):
        monkeypatch.delenv(key,raising=False)
    monkeypatch.setenv('AI_PROVIDER','demo');monkeypatch.setenv('IMAGE_PROVIDER','none')
    return create_app({'TESTING':True,'SECRET_KEY':'remaining-tests','DATA_DIR':str(tmp_path)}).test_client()

def workflow(client):
    script='INT. OFFICE - DAY\nMira, wearing a blue jacket, opens her laptop.\nMira: Ready?\nPROCESS: 13 ECU laptop workflow shots, static tripod, same screen setup.\nEXT. STREET - NIGHT\nMira walks to the gate.'
    plan=fixture_plan(client,script)
    approval=post(client,'/api/approve',{'plan':plan,'reviewed':True}).json
    rows=[]
    for scene in dict.fromkeys(r['scene'] for r in plan):
        response=post(client,'/api/breakdown',{'plan':plan,'approval':approval['approval'],'scene':scene})
        assert response.status_code==200,response.json
        rows.extend(response.json['shots'])
    return {'plan':plan,'breakdown':rows,'approval':approval['approval'],'breakdown_hash':approval['plan_hash'],'visual_context':'Pencil storyboard','continuity':{'characters':{'Mira':'Adult, short curly hair'},'locations':{'OFFICE':'Oak desk and north window'},'scenes':{'1':'Blue jacket, daylight'}}}

def png():
    out=BytesIO();Image.new('RGB',(64,36),'#71827b').save(out,format='PNG');return out.getvalue()

def upload(client,data,index=0):
    result=post(client,'/api/storyboard/upload',{**data,'shot_id':data['breakdown'][index]['id'],'image':base64.b64encode(png()).decode()})
    assert result.status_code==200,result.json
    return result.json['image']

def approve(client,data,image):
    r=post(client,'/api/storyboard/approve',{**data,'shot_id':image['shot_id'],'asset_id':image['asset_id'],'approved':True})
    assert r.status_code==200,r.json
    return r.json['image']

class Reply:
    def __init__(self,data):self.data=data if isinstance(data,bytes) else json.dumps(data).encode()
    def __enter__(self):return self
    def __exit__(self,*args):pass
    def read(self,n):return self.data[:n]

def cloud(monkeypatch):
    monkeypatch.setenv('AI_PROVIDER','openai');monkeypatch.setenv('OPENAI_API_KEY','test-key-not-real');monkeypatch.setenv('OPENAI_MODEL','test-text-model')

def image_cloud(monkeypatch):
    monkeypatch.setenv('IMAGE_PROVIDER','openai');monkeypatch.setenv('IMAGE_API_KEY','test-image-key');monkeypatch.setenv('IMAGE_MODEL','test-image-model')

def test_cloud_text_round_trip_and_source_protection(client,monkeypatch):
    cloud(monkeypatch);calls=[]
    def remote(req,timeout):
        payload=json.loads(req.data);calls.append((req.full_url,payload,timeout))
        inputs=json.loads(payload['messages'][1]['content'])
        rows=copy.deepcopy(inputs.get('draft',inputs.get('conservative_draft')))
        if 'conservative_draft' in inputs:
            for r in rows:r['notes']='AI proposed setup';r['script']='INVENTED DIALOGUE'
        else:
            # Provider returns only its editable string fields and source identities.
            rows=[{k:r[k] for k in ('beat','source_ids','shot_description','shot_size','shot_angle','cast')} for r in rows]
        return Reply({'choices':[{'message':{'content':json.dumps({'shots':rows})}}]})
    monkeypatch.setattr(providers,'urlopen',remote)
    data=workflow(client)
    assert calls and all(c[0].endswith('/chat/completions') for c in calls)
    assert all(c[1]['response_format']=={'type':'json_object'} for c in calls)
    assert any(r['coverage_label'] for r in data['breakdown'])
    assert all(r['source_spans'] for r in data['plan'])
    assert all('INVENTED' not in r['script'] for r in data['breakdown'])
    assert client.get('/api/config').json['generation_mode']=='live_ai'

@pytest.mark.parametrize('outer',[[],{},None,{'choices':[]},{'choices':[{'message':{'content':None}}]},{'choices':[{'message':{'content':'{"shots":'}}]}])
def test_cloud_bad_envelopes(client,monkeypatch,outer):
    cloud(monkeypatch);monkeypatch.setattr(providers,'urlopen',lambda *a,**k:Reply(outer))
    with pytest.raises(UserError) as error:providers.complete('test',{})
    assert error.value.status==502

@pytest.mark.parametrize('code,status',[(401,503),(403,503),(429,503),(500,502)])
def test_provider_http_errors_redacted(monkeypatch,code,status):
    def fail(*a,**kw):raise HTTPError('https://example.test',code,'SECRET PROVIDER BODY',{},None)
    monkeypatch.setattr(providers,'urlopen',fail)
    with pytest.raises(UserError) as error:providers.post_json('https://example.test',{})
    assert error.value.status==status and 'SECRET' not in str(error.value)

@pytest.mark.parametrize('error,status',[(TimeoutError(),504),(URLError('secret'),503)])
def test_provider_network_errors(monkeypatch,error,status):
    def fail(*a,**k):raise error
    monkeypatch.setattr(providers,'urlopen',fail)
    with pytest.raises(UserError) as caught:providers.post_json('https://example.test',{})
    assert caught.value.status==status

@pytest.mark.parametrize('raw',[b'not json',b'\xff',b'x'*16_000_001],ids=['not-json','invalid-utf8','oversize'])
def test_provider_bad_wire_data(monkeypatch,raw):
    monkeypatch.setattr(providers,'urlopen',lambda *a,**k:Reply(raw))
    with pytest.raises(UserError) as error:providers.post_json('https://example.test',{})
    assert error.value.status==502

def test_ollama_preserved_and_envelope_validated(client,monkeypatch):
    monkeypatch.setenv('AI_PROVIDER','ollama');calls=[]
    monkeypatch.setattr(providers,'urlopen',lambda req,timeout:(calls.append(req) or Reply({'message':{'content':'{"shots":[]}'}})))
    assert providers.complete('test',{})=={'shots':[]}
    assert calls[0].full_url.endswith('/api/chat')
    monkeypatch.setattr(providers,'urlopen',lambda *a,**k:Reply({'message':[]}))
    with pytest.raises(UserError):providers.complete('test',{})

def test_no_image_provider_never_fakes_generation(client,monkeypatch):
    data=workflow(client);monkeypatch.setenv('OPENAI_API_KEY','text-key-only')
    assert not client.get('/api/config').json['image_enabled']
    r=post(client,'/api/storyboard/generate',{**data,'shot_id':data['plan'][0]['id']})
    assert r.status_code==503 and 'not configured' in r.json['error']
    with client.application.extensions['store'].connect() as db:assert db.execute('SELECT COUNT(*) FROM assets').fetchone()[0]==0

def test_continuity_contains_every_camera_and_scene_input(client):
    data=workflow(client);row=data['breakdown'][0]
    ctx=storyboard.context(data['breakdown'],row,data['visual_context'],data['continuity'])
    assert ctx['characters']['Mira']=='Adult, short curly hair'
    assert ctx['location']=='Oak desk and north window'
    assert ctx['scene']=='Blue jacket, daylight'
    assert set(storyboard.VISUAL_FIELDS)==set(ctx['shot'])
    assert len(ctx['scene_shots'])>=2
    assert all(s['shot_location']=='OFFICE' for s in ctx['scene_shots'])

def test_upload_replace_approve_and_no_automatic_approval(client):
    data=workflow(client);first=upload(client,data)
    assert first['source']=='uploaded' and not first['approved']
    assert first['data'].startswith('data:image/jpeg;base64,')
    approved=approve(client,data,first);assert approved['approved']
    replacement=upload(client,data)
    assert replacement['asset_id']!=first['asset_id'] and not replacement['approved']

@pytest.mark.parametrize('change',['shot_description','lighting','costumes','cast','shot_location','size_angle','camera_movement','number','style','continuity'])
def test_outdated_image_cannot_be_approved_or_exported(client,change):
    data=workflow(client);image=approve(client,data,upload(client,data))
    data['images']={image['shot_id']:image}
    if change=='style':data['visual_context']='Changed style'
    elif change=='continuity':data['continuity']['characters']['Mira']='Different appearance'
    elif change=='number':data['breakdown'][0]['number']='100'
    else:data['breakdown'][0][change]='Changed'
    r=post(client,'/api/storyboard/approve',{**data,'shot_id':image['shot_id'],'asset_id':image['asset_id'],'approved':True})
    assert r.status_code==409,r.json
    r=post(client,'/api/export/pptx',data);assert r.status_code==409,r.json

@pytest.mark.parametrize('per_slide',[1,2])
@pytest.mark.parametrize('use_image',[False,True])
def test_ppt_with_and_without_approved_image(client,per_slide,use_image):
    data=workflow(client);data['per_slide']=per_slide
    if use_image:
        image=approve(client,data,upload(client,data,1));data['images']={image['shot_id']:image}
    r=post(client,'/api/export/pptx',data);assert r.status_code==200,r.json
    deck=Presentation(BytesIO(r.data));assert len(deck.slides)==(len(data['plan'])+per_slide-1)//per_slide
    pictures=[(i,s) for i,slide in enumerate(deck.slides) for s in slide.shapes if s.shape_type==13]
    assert len(pictures)==int(use_image)
    if use_image:assert pictures[0][0]==1//per_slide
    for i,row in enumerate(data['breakdown']):
        slide=deck.slides[i//per_slide]
        assert 'SHOT '+row['number'] in ' '.join(s.text for s in slide.shapes if s.has_text_frame)
        assert row['source_script_block'] in slide.notes_slide.notes_text_frame.text

def test_unapproved_and_forged_images_never_enter_ppt(client):
    data=workflow(client);image=upload(client,data);image['approved']=True
    r=post(client,'/api/export/pptx',{**data,'images':{image['shot_id']:image}})
    assert r.status_code==409

def test_optional_approved_reference_in_excel(client):
    from openpyxl import load_workbook
    data=workflow(client)
    r=post(client,'/api/export/xlsx',data);assert r.status_code==200,r.json
    assert not load_workbook(BytesIO(r.data)).active._images
    image=approve(client,data,upload(client,data,1));data['images']={image['shot_id']:image}
    r=post(client,'/api/export/xlsx',data);assert r.status_code==200,r.json
    ws=load_workbook(BytesIO(r.data)).active
    assert len(ws._images)==1 and ws._images[0].anchor._from.col==2 and ws._images[0].anchor._from.row==7

def test_image_generation_and_reference_edit_transport(client,monkeypatch):
    data=workflow(client);image_cloud(monkeypatch);requests=[]
    def remote(req,timeout):
        requests.append(req)
        return Reply({'data':[{'b64_json':base64.b64encode(png()).decode()}]})
    monkeypatch.setattr(providers,'urlopen',remote)
    r=post(client,'/api/storyboard/generate',{**data,'shot_id':data['plan'][0]['id']});assert r.status_code==200,r.json
    image=approve(client,data,r.json['image']);data['images']={image['shot_id']:image}
    r=post(client,'/api/storyboard/generate',{**data,'shot_id':data['plan'][1]['id']});assert r.status_code==200,r.json
    assert requests[0].full_url.endswith('/images/generations')
    assert requests[1].full_url.endswith('/images/edits')
    assert b'name="image[]"' in requests[1].data
    assert b'Adult, short curly hair' in requests[0].data
    assert r.json['image']['reference_count']==1 and not r.json['image']['approved']

@pytest.mark.parametrize('response',[{},None,{'data':[]},{'data':[{'b64_json':'garbage'}]},{'data':[{'url':'https://untrusted.test/image'}]}])
def test_invalid_generated_image_rejected(client,monkeypatch,response):
    data=workflow(client);image_cloud(monkeypatch)
    monkeypatch.setattr(providers,'urlopen',lambda *a,**k:Reply(response))
    r=post(client,'/api/storyboard/generate',{**data,'shot_id':data['plan'][0]['id']});assert r.status_code==502

def test_grouped_range_uses_one_representative_image(client):
    data=workflow(client);index=next(i for i,r in enumerate(data['breakdown']) if r['coverage_label'])
    image=approve(client,data,upload(client,data,index));row=data['breakdown'][index]
    assert image['number']==row['number'] and image['representative_shot']==row['representative_shot']
    result=post(client,'/api/export/pptx',{**data,'images':{image['shot_id']:image}})
    deck=Presentation(BytesIO(result.data))
    assert len(deck.slides)==len(data['breakdown'])
    assert sum(s.shape_type==13 for slide in deck.slides for s in slide.shapes)==1

@pytest.mark.parametrize('bad',['not base64','data:image/svg+xml;base64,PHN2Zz48L3N2Zz4=',None])
def test_upload_rejects_invalid_content(client,bad):
    data=workflow(client)
    r=post(client,'/api/storyboard/upload',{**data,'shot_id':data['plan'][0]['id'],'image':bad});assert r.status_code==400

def test_project_and_image_survive_app_restart(client):
    data=workflow(client);image=approve(client,data,upload(client,data))
    project={'version':3,'title':'Durable project','metadata':{'director':'Mira'},'script':'Original script','plan':data['plan'],'breakdown':data['breakdown'],'approval':data['approval'],'breakdownHash':data['breakdown_hash'],'images':{image['shot_id']:image},'continuity':data['continuity'],'visualContext':data['visual_context']}
    saved=post(client,'/api/projects',{'project':project});assert saved.status_code==200
    new=create_app({'TESTING':True,'SECRET_KEY':'remaining-tests','DATA_DIR':client.application.config['DATA_DIR']}).test_client()
    assert len(new.get('/api/projects').json['projects'])==1
    restored=new.get('/api/projects/'+saved.json['id']).json['project']
    assert restored['images'][image['shot_id']]['approved'] and restored['metadata']['director']=='Mira'
    assert restored['continuity']==data['continuity']
    result=post(new,'/api/export/pptx',{**data,'images':restored['images']});assert result.status_code==200
    assert any(s.shape_type==13 for slide in Presentation(BytesIO(result.data)).slides for s in slide.shapes)

def test_cross_origin_login_and_authenticated_requests(client,monkeypatch,tmp_path):
    monkeypatch.setenv('APP_ENV','production');monkeypatch.setenv('SECRET_KEY','production-test-secret')
    monkeypatch.setenv('APP_ACCESS_PASSWORD','workspace-test-password');monkeypatch.setenv('DATA_DIR',str(tmp_path));monkeypatch.setenv('FRONTEND_ORIGINS','https://studio.example')
    c=create_app({'TESTING':True}).test_client();origin={'Origin':'https://studio.example','X-Script2Set':'1'}
    pre=c.options('/api/projects',headers={**origin,'Access-Control-Request-Headers':'Authorization,Content-Type,X-Script2Set'})
    assert pre.status_code==204 and pre.headers['Access-Control-Allow-Origin']=='https://studio.example'
    assert c.get('/api/projects',headers=origin).status_code==401
    login=c.post('/api/login',json={'password':'workspace-test-password'},headers=origin)
    token=login.json['token'];other=create_app({'TESTING':True}).test_client()
    assert other.get('/api/projects',headers={**origin,'Authorization':'Bearer '+token}).status_code==200
    denied=other.get('/api/projects',headers={'Origin':'https://evil.example','Authorization':'Bearer '+token})
    assert denied.status_code==403 and 'Access-Control-Allow-Origin' not in denied.headers
    assert other.get('/api/projects',headers={**origin,'Authorization':'Bearer forged'}).status_code==401

@pytest.mark.parametrize('origin',['*','http://insecure.example','https://example.test/path','https://user:pass@example.test'])
def test_production_cors_configuration_rejected(client,monkeypatch,tmp_path,origin):
    for k,v in {'APP_ENV':'production','SECRET_KEY':'secret','APP_ACCESS_PASSWORD':'password','DATA_DIR':str(tmp_path),'FRONTEND_ORIGINS':origin}.items():monkeypatch.setenv(k,v)
    assert create_app({'TESTING':True}).test_client().get('/api/config').status_code==503

def test_frontend_build_has_no_secrets_and_no_localhost(client):
    root=Path(__file__).resolve().parents[1]
    env={**os.environ,'PUBLIC_API_BASE_URL':'https://api.example.test','OPENAI_API_KEY':'DO-NOT-PUBLISH-KEY'}
    result=subprocess.run([shutil.which('node'),'scripts/build_frontend.mjs'],cwd=root,env=env,capture_output=True,text=True)
    assert result.returncode==0,result.stderr
    out=root/'frontend-dist';assert '/print /print.html 200' in (out/'_redirects').read_text()
    assert 'https://api.example.test' in (out/'static/config.js').read_text()
    assert 'connect-src' in (out/'_headers').read_text()
    for file in out.rglob('*'):
        if file.is_file():assert 'DO-NOT-PUBLISH-KEY' not in file.read_text(encoding='utf-8')
    env['PUBLIC_API_BASE_URL']='http://localhost:5000'
    assert subprocess.run([shutil.which('node'),'scripts/build_frontend.mjs'],cwd=root,env=env,capture_output=True).returncode!=0
