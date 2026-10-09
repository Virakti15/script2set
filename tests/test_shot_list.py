from io import BytesIO, StringIO
import csv
import json
from datetime import datetime
import pytest
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from services.schema import normalize, validate_plan, production_view, ordered, UserError, PRODUCTION
from services.exporting import excel_export, csv_export, ppt_export
from services.importing import import_excel
from services.screenplay import analyze
from services.planning import generate_plan
from test_architecture import planned, broken_down
from test_workflow import client, post

@pytest.mark.parametrize('number',['1','7-12','29–41 Process','42—43','131-140','1.2.3'])
def test_shot_text_survives_excel_csv_and_ppt(number):
    row=normalize({'number':number,'script':'Phone screen','shot_description':'ECU of phone'})
    validate_plan([row]);sheet=load_workbook(excel_export([row])).active
    assert sheet['A7'].value==number and sheet['A7'].data_type=='s' and sheet['A7'].number_format=='@'
    imported=import_excel(excel_export([row]).getvalue())['shots'][0]
    assert imported['number']==number
    assert list(csv.DictReader(StringIO(csv_export([row]).getvalue().decode('utf-8-sig'))))[0]['SHOT']==number
    assert number in ' '.join(s.text for s in Presentation(ppt_export([row])).slides[0].shapes if s.has_text_frame)

@pytest.mark.parametrize('number',['12-7','7-foo','7/12','=1+1'])
def test_malformed_ranges_rejected(number):
    with pytest.raises(UserError):validate_plan([{'number':number,'script':'x'}])

def test_overlapping_ranges_rejected():
    with pytest.raises(UserError,match='overlap'):validate_plan([{'number':'7-12','script':'x'},{'number':'12-14','script':'y'}])

def test_production_order_wins_over_number_labels():
    rows=[normalize({'number':n,'script':'x','production_order':i}) for i,n in enumerate(['3','22','29–41 Process','42–43','56–85','87–107'])]
    assert [r['number'] for r in ordered(list(reversed(rows)))]==[r['number'] for r in rows]

def test_process_count_group_not_thirteen_repetitive_rows():
    _,rows=planned('INT. OFFICE - DAY\nPROCESS: 13 ECU laptop workflow shots, static tripod, same screen setup.\nREACTION: Mira smiles.')
    assert len(rows)==2 and rows[0]['number']=='1–13 Process' and rows[1]['number']=='14'
    assert rows[0]['coverage_count']==13 and rows[0]['representative_shot']=='1'

def test_two_phone_pov_shots_share_one_row():
    _,rows=planned('INT. PHONE SETUP - DAY\nPhone POV x2: two shots of the same phone interface, locked off.')
    assert len(rows)==1 and rows[0]['number']=='1–2 Coverage'

def test_unspecified_broll_count_not_invented():
    _,_,rows=broken_down('EXT. HILL ROAD - DAY\nB-ROLL: Shots of people, lanes and chaos, same street coverage setup.')
    assert len(rows)==1 and rows[0]['number']=='1 B-roll' and rows[0]['coverage_count']==0
    assert 'count to confirm' in rows[0]['assumptions']

def test_different_locations_or_setups_never_group():
    _,rows=planned('EXT. HILL ROAD - DAY\nB-ROLL: 4 wide street shots.\nINT. SHOP - NIGHT\nB-ROLL: 3 ECU clothing shots.')
    assert len(rows)==2 and rows[0]['scene']!=rows[1]['scene']
    assert rows[0]['number']=='1–4 B-roll' and rows[1]['number']=='5–7 B-roll'

def test_match_cut_motion_blur_long_take_notes():
    _,_,rows=broken_down('INT. MEDITATION AREA - DAY\nECU of Mira eye, frontal, static tripod, 24 FPS.\nMATCH CUT: between spherical objects; motion blur.\nLONG TAKE: Mira walks across the room.')
    assert 'MATCH CUT' in rows[0]['notes'] and 'motion blur' in rows[0]['notes']
    assert 'LONG TAKE' in rows[1]['notes']
    assert rows[0]['equipment']=='Tripod' and rows[0]['fps']=='24 FPS'

def test_graphic_projection_and_optional_blanks():
    _,_,rows=broken_down('STAGE 2: ANGER');row=production_view(rows[0])
    assert row['size_angle']=='Full Frame Graphic' and row['shot_location']=='N/A'
    assert row['equipment']=='N/A' and row['cast']=='None' and row['props']=='None'
    assert row['clip_no']=='' and row['script']=='STAGE 2: ANGER'
    assert len(row)==19 and 'source_script_block' not in row

def test_metadata_once_and_working_table():
    row=normalize({'number':'29–41 Process','script':'Laptop workflow'})
    sheet=load_workbook(excel_export([row],{'client':'Example client','director':'Example director','aspect_ratio':'16:9'})).active
    assert sheet['A1'].value=='SHOT LIST' and sheet['B2'].value=='Example client'
    assert sheet['B3'].value=='Example director' and sheet['F4'].value=='16:9'
    assert sheet.freeze_panes=='D7' and sheet.auto_filter.ref=='A6:S7' and all(r.max_row<6 for r in sheet.merged_cells.ranges)
    imported=import_excel(excel_export([row],{'client':'Example client'}).getvalue())
    assert imported['metadata']['client']=='Example client'

def test_reference_style_thirteen_column_import():
    wb=Workbook();sheet=wb.active
    sheet.append(['SHOT LIST']);sheet.append(['Director','Example director']);sheet.append(list(PRODUCTION.values())[:13])
    for number,description in [('3','ECU of character eye'),('22','ECU of phone interface'),('29–41 Process','Multiple ECU laptop shots'),('42–43','Phone POV x2'),('56–85','B-roll people lanes chaos'),('87–107','Vendors clothes accessories')]:
        sheet.append([number,'DAY / INT','Phone / laptop','ECU / Top','Static','Phone setup','Tripod','24 FPS',description,'Daylight + Bounce','Narrator (V.O.): Make it happen.','Match cut',''])
    data=BytesIO();wb.save(data);result=import_excel(data.getvalue())
    assert len(result['shots'])==6 and result['metadata']['director']=='Example director'
    validate_plan(result['shots'])
    for row in result['shots']:
        assert row['shot_angle']=='Top' and row['equipment']=='Tripod' and row['existing']['notes']=='Match cut'

def test_corrupt_excel_dates_rejected():
    wb=Workbook();ws=wb.active;ws.append(['SHOT','SHOT DESCRIPTION']);ws.append([datetime(2026,7,12),'phone'])
    data=BytesIO();wb.save(data)
    with pytest.raises(UserError,match='Excel date'):import_excel(data.getvalue())

def test_breakdown_edit_range_notes_and_framing_reaches_all_exports(client):
    _,plan=planned('INT. OFFICE - DAY\nMira: Hello.')
    approval=post(client,'/api/approve',{'plan':plan,'reviewed':True}).json
    rows=post(client,'/api/breakdown',{'plan':plan,'approval':approval['approval'],'scene':'1'}).json['shots']
    rows[0].update(number='29–41 Process',size_angle='ECU / Top',shot_description='ECU of Mira speaking',equipment='Magic Arm',notes='Manual match-cut note',clip_no='A001')
    data={'plan':plan,'breakdown':rows,'approval':approval['approval'],'breakdown_hash':approval['plan_hash'],'metadata':{'director':'Review director'}}
    printed=post(client,'/api/print',data).json
    assert printed['shots'][0]['number']=='29–41 Process' and printed['shots'][0]['shot_angle']=='Top'
    assert printed['metadata']['director']=='Review director'
    for kind in ('xlsx','csv','pptx'):assert post(client,'/api/export/'+kind,data).status_code==200

def test_old_30_columns_migrate_without_losing_special_notes(client):
    old={'id':'fixed','number':'1.1.1','scene':'1','beat':'1.1','script':'x','sfx':'Bell','sound':'Jazz','vfx':'Remove wire','continuity':'Same jacket','art_setup':'Office dressing'}
    p=post(client,'/api/migrate',{'project':{'version':2,'plan':[old],'breakdown':[old]}}).json['project']
    assert p['version']==3 and p['stale']
    notes=p['breakdown'][0]['notes']
    assert all(v in notes for v in ('Bell','Jazz','Remove wire','Same jacket'))
    assert p['plan'][0]['number']=='1.1.1'

def test_model_omitted_source_still_rejected(monkeypatch):
    scene=analyze('INT. ROOM - DAY\nPROCESS: 3 ECU laptop workflow shots.\nREACTION: Mira smiles.')['chunks'][0]
    first=scene['beats'][0];monkeypatch.setenv('AI_PROVIDER','ollama')
    proposed={'beat':first['number'],'source_ids':first['source_ids'],'shot_description':'Laptop workflow','shot_size':'ECU'}
    monkeypatch.setattr('services.planning.complete',lambda *args:{'shots':[proposed]})
    with pytest.raises(UserError,match='omitted'):generate_plan(scene)

def test_lighting_notes_do_not_displace_dialogue_description():
    _,_,rows=broken_down('INT. OFFICE - DAY\nLIGHTING: daylight bounce.\nMira: Ready?\nRavi: Yes.')
    assert len(rows)==1 and 'exchange' in rows[0]['shot_description']
    assert 'LIGHTING:' not in rows[0]['shot_description']

def test_explicit_no_talent_group_and_phone_pov_framing():
    _,_,rows=broken_down('EXT. MARKET - DAY\nB-ROLL: 5 ECU vendor clothes shots. No talent in frame.\nPhone POV x2: same phone interface, no hands in frame.')
    assert rows[0]['cast']=='None' and rows[0]['costumes']=='N/A'
    assert rows[1]['shot_size']=='Close-Up' and rows[1]['cast']=='None'

def test_cleared_optional_notes_stay_cleared_in_exports():
    row=normalize({'number':'1','script':'x','notes':'','art_graphics':'','sfx':'old sound','art_setup':'old art'})
    view=production_view(row)
    assert view['notes']=='' and view['art_graphics']==''

def test_edited_range_updates_representative_shot(client):
    _,plan=planned('Mira: Hello');approval=post(client,'/api/approve',{'plan':plan,'reviewed':True}).json
    rows=post(client,'/api/breakdown',{'plan':plan,'approval':approval['approval'],'scene':'1'}).json['shots'];rows[0]['number']='29-41 Process'
    data={'plan':plan,'approval':approval['approval'],'breakdown_hash':approval['plan_hash'],'breakdown':rows}
    row=post(client,'/api/print',data).json['shots'][0]
    assert row['representative_shot']=='29' and row['coverage_count']==13
