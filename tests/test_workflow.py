from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import csv
import json
import pytest
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from app import create_app
from services.schema import normalize, ordered, UserError
from services.screenplay import analyze
from services.planning import generate_plan
from services.importing import import_excel
from services.exporting import ppt_export, csv_export

@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv('AI_PROVIDER','demo')
    monkeypatch.delenv('APP_ACCESS_PASSWORD',raising=False)
    monkeypatch.delenv('VERCEL',raising=False)
    monkeypatch.delenv('APP_ENV',raising=False)
    return create_app({'TESTING':True,'SECRET_KEY':'test-only-secret'}).test_client()

def post(client,path,data):
    return client.post(path,json=data,headers={'X-Script2Set':'1'})

def fixture_plan(client,script=None):
    script=script or Path('static/sample.txt').read_text(encoding='utf-8')
    analysis=post(client,'/api/analyze',{'script':script})
    assert analysis.status_code==200,analysis.json
    plan=[]
    for chunk in analysis.json['chunks']:
        result=post(client,'/api/plan',{'scene':chunk,'offset':len(plan),'previous':plan[-4:]})
        assert result.status_code==200,result.json
        plan+=result.json['shots']
    return plan

def completed(client):
    plan=fixture_plan(client)
    approval=post(client,'/api/approve',{'plan':plan,'reviewed':True}).json
    result=post(client,'/api/breakdown',{'plan':plan,'approval':approval['approval'],'scene':'1'})
    assert result.status_code==200,result.json
    return {'plan':plan,'approval':approval['approval'],'breakdown':result.json['shots'],'breakdown_hash':approval['plan_hash']}

def test_reference_script_retained(client):
    plan=fixture_plan(client)
    assert plan[-1]['script']=='STAGE 6: DELEGATION'
    cards=[r for r in plan if r['shot_type']=='Title Card / Graphic']
    assert len(cards)==7
    assert all(r['cast']=='None' and r['props']=='None' and r['shot_size']=='Full Screen Graphic' for r in cards)
    assert len([r for r in plan if 'What happened to her?' in r['script']])==1
    opening=plan[0]
    assert 'What?' in opening['script'] and 'Maanik' in opening['cast']
    assert not any(r['script'].strip() in ('"','Aarti:') for r in plan)
    assert not any('tiger' in r['props'].lower() for r in plan)

@pytest.mark.parametrize('source', ['', '  ', '"\n"', 'Aarti:', 'INT. OFFICE - DAY'])
def test_empty_or_artifact_script(client,source):
    assert post(client,'/api/analyze',{'script':source}).status_code==400

def test_broken_quotes(client):
    plan=fixture_plan(client,'"\n"Aarti:\nHello"\n\nTITLE CARD: END')
    assert plan[0]['script']=='Aarti: Hello'
    assert len(plan)==2

@pytest.mark.parametrize('source,expected', [
    ('TITLE CARD: Hello','Title Card'),('STAGE 6: END','Stage Card'),('SUPER: Monday','On-Screen Text'),
    ('SFX: crash','SFX'),('MUSIC: jazz','Music Cue'),('MONTAGE: city','Montage'),
    ('CAMERA tracks left','Camera Direction'),('Narrator (V.O.): Hello','Voice Over'),
    ('Aarti: Hello','Dialogue'),('Aarti slams the table.','Action')])
def test_classification(source,expected):
    assert analyze(source)['elements'][0]['type']==expected

def test_scene_memory(client):
    plan=fixture_plan(client,'INT. OFFICE - DAY\nAarti: Hello\n\nEXT. STREET - NIGHT\nMaanik: Goodbye')
    assert plan[0]['shot_location']=='OFFICE'
    assert plan[1]['shot_location']=='STREET'
    assert plan[1]['day_night']=='NIGHT'

def test_long_script_chunks():
    data=analyze('\n\n'.join('Aarti: '+('Hello. '*80) for _ in range(50)))
    assert len(data['chunks'])>1
    assert len(data['elements'])==50

def test_approval_required(client):
    plan=fixture_plan(client)
    assert post(client,'/api/approve',{'plan':plan}).status_code==409
    assert post(client,'/api/breakdown',{'plan':plan,'scene':'1'}).status_code==409

@pytest.mark.parametrize('kind',['csv','xlsx','pptx'])
def test_real_exports_latest_edits(client,kind):
    data=completed(client);data['breakdown'][0]['props']='Edited brass lamp'
    response=post(client,'/api/export/'+kind,data)
    assert response.status_code==200,response.json
    assert 'attachment' in response.headers['Content-Disposition']
    if kind=='csv':
        assert response.data.startswith(b'\xef\xbb\xbf')
        assert 'Edited brass lamp' in response.data.decode('utf-8-sig')
    elif kind=='xlsx':
        wb=load_workbook(BytesIO(response.data));assert wb.active.max_column==19
        assert 'Edited brass lamp' in list(wb.active.values)[6]
    else:
        deck=Presentation(BytesIO(response.data));assert len(deck.slides)==len(data['plan'])
        assert 'Edited brass lamp' in deck.slides[0].notes_slide.notes_text_frame.text
        assert any('PENDING' in s.text for s in deck.slides[0].shapes if s.has_text_frame)

@pytest.mark.parametrize('route',['/api/export/csv','/api/export/xlsx','/api/export/pptx','/api/print','/api/breakdown'])
def test_stale_plan_rejected_server_side(client,route):
    data=completed(client);data['plan'][0]['shot_description']='Changed after approval'
    response=post(client,route,data)
    assert response.status_code==409
    assert 'Shot plan changed' in response.json['error']

def test_empty_export(client):
    assert post(client,'/api/export/pptx',{'plan':[]}).status_code==400

def test_ppt_two_up(client):
    data=completed(client);data['per_slide']=2
    response=post(client,'/api/export/pptx',data)
    deck=Presentation(BytesIO(response.data))
    assert len(deck.slides)==(len(data['plan'])+1)//2
    assert all(any(s.has_text_frame and 'SHOT' in s.text for s in slide.shapes) for slide in deck.slides)

def test_numeric_sort_and_legacy_schema():
    rows=[{'shot_number':'1.'+str(i),'scene_number':1,'visual_action':'Action','suggested_shot_size':'Wide'} for i in [10,2,11,1]]
    rows=[normalize(r) for r in rows]
    assert [r['number'] for r in ordered(rows)]==['1.1','1.2','1.10','1.11']
    deck=Presentation(ppt_export(rows))
    assert 'Wide' in ' '.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)

def excel_fixture():
    wb=Workbook();wb.active.title='Read me';wb.active.append(['Not the shot sheet'])
    ws=wb.create_sheet('Shot list');ws.append(['Production: sample']);ws.append(['Shot Number','Script','Cast','Props','Costumes','Shot Location','FPS'])
    ws.append(['1.1','Aarti: Hello','Aarti','Red phone','Blue jacket','Studio A','25'])
    out=BytesIO();wb.save(out);return out.getvalue()

def test_excel_preserves_user_fields(client):
    data=import_excel(excel_fixture());assert data['worksheet']=='Shot list'
    plan=data['shots'];approval=post(client,'/api/approve',{'plan':plan,'reviewed':True}).json
    response=post(client,'/api/breakdown',{'plan':plan,'approval':approval['approval'],'scene':'1'})
    row=response.json['shots'][0]
    assert row['costumes']=='Blue jacket' and row['props']=='Red phone' and row['fps']=='25'
    assert row['shot_location']=='Studio A'

def test_excel_upload_endpoint(client):
    response=client.post('/api/import',data={'file':(BytesIO(excel_fixture()),'shots.xlsx')},headers={'X-Script2Set':'1'})
    assert response.status_code==200

@pytest.mark.parametrize('data',[b'not a zip',b'',b'PK123'])
def test_invalid_excel(data):
    with pytest.raises(UserError):import_excel(data)

@pytest.mark.parametrize('formula',['=1+1','+cmd','-cmd','@SUM(A1)','  =HYPERLINK("bad")'])
def test_csv_formula_injection(formula):
    csv_data=csv_export([{'number':'1.1','script':formula}]).getvalue().decode('utf-8-sig')
    records=list(csv.DictReader(csv_data.splitlines()))
    assert records[0]['SCRIPT / DIALOGUE'].startswith("'")
    assert 'Source Script Block' not in records[0]

def test_origin_and_csrf(client):
    assert client.post('/api/analyze',json={'script':'Hello'}).status_code==403
    assert client.post('/api/analyze',json={'script':'Hello'},headers={'X-Script2Set':'1','Origin':'https://unrelated.example'}).status_code==403

def test_security_headers_and_health(client):
    r=client.get('/');assert 'frame-ancestors' in r.headers['Content-Security-Policy']
    assert client.get('/health').json=={'status':'ok','ai_provider':'demo','environment':'local'}

def test_forged_approval(client):
    data=completed(client);data['approval']='forged';assert post(client,'/api/export/csv',data).status_code==409

def test_unsupported_provider_is_clear(client,monkeypatch):
    monkeypatch.setenv('AI_PROVIDER','unknown')
    scene=analyze('Aarti: Hello')['chunks'][0]
    response=post(client,'/api/plan',{'scene':scene})
    assert response.status_code==503 and 'Unsupported AI provider' in response.json['error']

def test_missing_cloud_key(client,monkeypatch):
    monkeypatch.setenv('AI_PROVIDER','openai');monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    response=post(client,'/api/plan',{'scene':analyze('Aarti: Hello')['chunks'][0]})
    assert response.status_code==503 and 'not configured' in response.json['error']

def test_cloud_malformed_response(client,monkeypatch):
    monkeypatch.setenv('AI_PROVIDER','ollama')
    monkeypatch.setattr('services.providers.post_json',lambda *a,**k:{'message':{'content':'nonsense'}})
    response=post(client,'/api/plan',{'scene':analyze('Aarti: Hello')['chunks'][0]})
    assert response.status_code==502 and 'structured data' in response.json['error']

def test_ai_dropped_card_rejected(client,monkeypatch):
    monkeypatch.setenv('AI_PROVIDER','ollama')
    monkeypatch.setattr('services.providers.post_json',lambda *a,**k:{'message':{'content':json.dumps({'shots':[{'source_ids':[1],'shot_description':'Two shot','shot_size':'Wide'}]})}})
    response=post(client,'/api/plan',{'scene':analyze('Aarti: Hello\n\nSTAGE 6: END')['chunks'][0]})
    assert response.status_code==502 and 'omitted' in response.json['error']
