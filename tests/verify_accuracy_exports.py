"""Reproducible, offline route-level sample review; no claimed live provider calls."""
import os,json,csv,html
from pathlib import Path
from io import BytesIO,StringIO
from openpyxl import load_workbook
from pptx import Presentation
from app import create_app
from services.schema import PRODUCTION,production_view

SCRIPTS={
 'dialogue': '''INT. EDIT SUITE - NIGHT
Mira wears a navy jacket and sets a red folder on the desk.
LIGHTING: warm desk lamp.
Mira: Did you check the final cut?
Ravi: Twice. The last line is still missing.
Mira: Keep the pause. Then use the original recording.
Ravi: I'll put it back.
REACTION: Mira smiles with relief.
''',
 'process': '''INT. KITCHEN - DAY
LIGHTING: daylight from the side window and bounce.
MS of Mira placing a saucepan on the stove, static tripod, 24 FPS.
PROCESS: 6 ECU shots of tea preparation in the same saucepan setup, top angle, static tripod, 24 FPS. Show water, tea leaves, stirring, simmering, straining and pouring. No talent in frame.
NARRATOR (V.O.): Six small steps. One quiet moment.
CU of Mira lifting the cup, eye level, static tripod, 24 FPS.
''',
 'multi-scene': '''INT. OFFICE - DAY
Mira picks up a red folder.
Mira: I'll bring the original.
EXT. STATION - DUSK
WIDE: Mira walks toward the platform. Handheld camera.
NARRATOR (V.O.): The city changes as the light fades.
INT. EDIT SUITE - NIGHT
Mira puts the red folder on the desk.
Ravi: Just in time.
TITLE CARD: TO BE CONTINUED
'''
}

def main():
 os.environ['AI_PROVIDER']='demo';os.environ['IMAGE_PROVIDER']='none'
 root=Path(__file__).resolve().parents[1];out=root.parent/'excel-priority-review';out.mkdir(exist_ok=True)
 client=create_app({'TESTING':True,'SECRET_KEY':'accuracy-sample-only','DATA_DIR':str(out/'test-storage')}).test_client()
 def post(path,data):
  r=client.post(path,json=data,headers={'X-Script2Set':'1'});assert r.status_code==200,(path,r.status_code,r.json);return r
 reports={}
 for name,script in SCRIPTS.items():
  analysis=post('/api/analyze',{'script':script}).json;plan=[]
  for chunk in analysis['chunks']:plan+=post('/api/plan',{'scene':chunk,'characters':analysis['characters'],'offset':len(plan),'previous':plan[-4:]}).json['shots']
  approval=post('/api/approve',{'plan':plan,'reviewed':True}).json;rows=[]
  for scene in dict.fromkeys(r['scene'] for r in plan):rows+=post('/api/breakdown',{'plan':plan,'approval':approval['approval'],'scene':scene}).json['shots']
  metadata={'client':'Script2Set accuracy review','director':'Confirmation Required','dop':'Confirmation Required','location':'See scene setups','date':'Confirmation Required','duration':'Confirmation Required','platform':'Production review','aspect_ratio':'Confirmation Required','reference':'Supplied CHATGPT SHOT BREAKDOWN.xlsx layout'}
  data={'plan':plan,'breakdown':rows,'approval':approval['approval'],'breakdown_hash':approval['plan_hash'],'source_script':script,'metadata':metadata}
  report=post('/api/validate',data).json;assert report['ok'] and report['source_verified'],report
  for ext in ('xlsx','csv'):(out/f'{name}.{ext}').write_bytes(post('/api/export/'+ext,data).data)
  for layout in (1,2):
   raw=post('/api/export/pptx',{**data,'per_slide':layout}).data;(out/f'{name}-{layout}up.pptx').write_bytes(raw)
   deck=Presentation(BytesIO(raw));assert len(deck.slides)==(len(rows)+layout-1)//layout
   for i,row in enumerate(rows):assert 'SHOT '+row['number'] in ' '.join(s.text for s in deck.slides[i//layout].shapes if s.has_text_frame)
  printed=post('/api/print',data).json;assert len(printed['shots'])==len(rows)
  articles=''.join('<article><h2>Scene '+r['scene']+' / Shot '+r['number']+'</h2><dl>'+''.join('<dt>'+html.escape(label)+'</dt><dd>'+html.escape(str(r[k]))+'</dd>' for k,label in PRODUCTION.items())+'</dl></article>' for r in printed['shots'])
  (out/f'{name}-print.html').write_text('<!doctype html><meta charset="utf-8"><style>'+(root/'static/print.css').read_text()+'</style>'+articles,encoding='utf-8')
  ws=load_workbook(out/f'{name}.xlsx').active
  assert ws.max_column==19 and ws.max_row==len(rows)+6
  assert [c.value for c in ws[6]]==list(PRODUCTION.values())
  assert ws.freeze_panes=='D7' and ws.auto_filter.ref==f'A6:S{ws.max_row}'
  assert all(m.max_row<6 for m in ws.merged_cells.ranges)
  assert ws['B2'].value==metadata['client'] and ws['F4'].value==metadata['aspect_ratio']
  for i,row in enumerate(rows,7):
   assert [c.value for c in ws[i]]==[v or None for v in production_view(row).values()]
   assert ws.cell(i,1).data_type=='s' and ws.cell(i,1).number_format=='@'
   assert ws.cell(i,1).comment
   if any(w.get('row_id')==row['id'] or not w.get('row_id') for w in report['warnings']):assert ws.cell(i,19).comment
   assert all(c.alignment.wrap_text for c in ws[i]) and ws.row_dimensions[i].height>=96
  csvrows=list(csv.DictReader(StringIO((out/f'{name}.csv').read_text(encoding='utf-8-sig'))))
  assert len(csvrows)==len(rows)
  assert [list(r.values()) for r in csvrows]==[list(production_view(r).values()) for r in rows]
  report['shots']=[{k:r[k] for k in ['number','scene','shot_size','shot_angle','camera_movement','shot_description','script','source_script_block','props','costumes','assumptions']} for r in rows]
  reports[name]=report;(out/f'{name}-source.txt').write_text(script,encoding='utf-8')
 (out/'sample-validation.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({k:{f:r[f] for f in ['rows','scenes','groups','ok','source_verified']} for k,r in reports.items()},indent=2))

if __name__=='__main__':main()
