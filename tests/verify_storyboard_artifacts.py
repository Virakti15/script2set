"""Create inspectable end-to-end artifacts with one explicitly uploaded test frame."""
import base64
import csv
import html
import json
import os
import tempfile
import sys
from io import BytesIO, StringIO
from pathlib import Path
from PIL import Image, ImageDraw
from openpyxl import load_workbook
from pptx import Presentation
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import create_app
from services.schema import PRODUCTION, production_view


def main():
    os.environ['AI_PROVIDER']='demo'
    os.environ['IMAGE_PROVIDER']='none'
    for key in ('APP_ACCESS_PASSWORD','APP_ENV','VERCEL','FRONTEND_ORIGINS'):
        os.environ.pop(key,None)
    root=Path(__file__).resolve().parents[1]
    out=root.parent/'storyboard-review';out.mkdir(exist_ok=True)
    script=(root/'tests/fixtures/production-review.txt').read_text(encoding='utf-8')
    with tempfile.TemporaryDirectory(dir=out) as data_dir:
        client=create_app({'TESTING':True,'SECRET_KEY':'artifact-review','DATA_DIR':data_dir}).test_client()
        def post(path,data):
            response=client.post(path,json=data,headers={'X-Script2Set':'1'})
            assert response.status_code==200,(path,response.status_code,response.get_json(silent=True))
            return response
        analysis=post('/api/analyze',{'script':script}).json;plan=[]
        for chunk in analysis['chunks']:
            plan+=post('/api/plan',{'scene':chunk,'characters':analysis['characters'],'offset':len(plan),'previous':plan[-4:]}).json['shots']
        approval=post('/api/approve',{'plan':plan,'reviewed':True}).json;rows=[]
        for scene in dict.fromkeys(r['scene'] for r in plan):
            rows+=post('/api/breakdown',{'plan':plan,'approval':approval['approval'],'scene':scene}).json['shots']
        grouped=next(r for r in rows if r.get('coverage_label'))
        continuity={'characters':{},'locations':{},'scenes':{}}
        style='Production storyboard validation. Uploaded frame, not AI generated.'
        frame=Image.new('RGB',(1280,720),(222,224,218));draw=ImageDraw.Draw(frame)
        draw.rectangle((45,45,1235,675),outline=(31,58,49),width=12)
        draw.text((90,90),'UPLOADED VALIDATION FRAME',fill=(31,58,49))
        draw.text((90,140),'Representative grouped coverage image',fill=(31,58,49))
        image_file=BytesIO();frame.save(image_file,'JPEG',quality=90)
        image_data=base64.b64encode(image_file.getvalue()).decode()
        common={'plan':plan,'breakdown':rows,'approval':approval['approval'],'breakdown_hash':approval['plan_hash'],
                'title':'Script2Set storyboard validation','metadata':{'client':'Validation production','director':'Test director','aspect_ratio':'16:9'},
                'visual_context':style,'continuity':continuity}
        uploaded=post('/api/storyboard/upload',{**common,'shot_id':grouped['id'],'image':image_data}).json['image']
        approved=post('/api/storyboard/approve',{**common,'shot_id':grouped['id'],'asset_id':uploaded['asset_id'],'approved':True}).json['image']
        payload={**common,'images':{grouped['id']:{'asset_id':approved['asset_id'],'approved':True}}}
        artifacts={}
        for kind in ('xlsx','csv'):
            artifacts[kind]=post('/api/export/'+kind,payload).data
            (out/f'production-breakdown.{kind}').write_bytes(artifacts[kind])
        printed=post('/api/print',payload).json
        (out/'print-payload.json').write_text(json.dumps(printed,ensure_ascii=False,indent=2),encoding='utf-8')
        esc=lambda value:html.escape(str(value or ''))
        metadata='<dl>'+''.join(f'<dt>{esc(printed["metadata_fields"][key])}</dt><dd>{esc(value)}</dd>' for key,value in printed['metadata'].items())+'</dl>'
        articles=''.join(f'<article><h2>Scene {esc(row["scene"])} / Shot {esc(row["number"])}</h2><dl>'+''.join(f'<dt>{esc(label)}</dt><dd>{esc(row[key])}</dd>' for key,label in PRODUCTION.items())+'</dl></article>' for row in printed['shots'])
        css=(root/'static/print.css').read_text(encoding='utf-8')
        (out/'production-print.html').write_text(f'<!doctype html><html><head><meta charset="utf-8"><title>Script2Set print validation</title><style>{css}</style></head><body><h1>SHOT LIST</h1>{metadata}{articles}</body></html>',encoding='utf-8')
        for count,name in ((1,'storyboard-one-up.pptx'),(2,'storyboard-two-up.pptx')):
            artifacts[name]=post('/api/export/pptx',{**payload,'per_slide':count}).data
            (out/name).write_bytes(artifacts[name])
        wb=load_workbook(BytesIO(artifacts['xlsx']));ws=wb.active
        assert ws.max_column==19 and list(list(ws.values)[5])==list(PRODUCTION.values())
        for index,(row,values) in enumerate(zip(rows,list(ws.values)[6:]),7):
            assert list(values)==[production_view(row)[k] or None for k in PRODUCTION]
            assert ws.cell(index,1).data_type=='s' and ws.cell(index,1).number_format=='@'
        csv_rows=list(csv.DictReader(StringIO(artifacts['csv'].decode('utf-8-sig'))))
        assert len(csv_rows)==len(rows) and all(list(record.values())==list(production_view(row).values()) for record,row in zip(csv_rows,rows))
        deck_checks={}
        for count,name in ((1,'storyboard-one-up.pptx'),(2,'storyboard-two-up.pptx')):
            deck=Presentation(BytesIO(artifacts[name]));pictures=[]
            assert len(deck.slides)==(len(rows)+count-1)//count
            for i,row in enumerate(rows):
                slide=deck.slides[i//count]
                assert 'SHOT '+row['number'] in ' '.join(shape.text for shape in slide.shapes if shape.has_text_frame)
                assert row['source_script_block'] in slide.notes_slide.notes_text_frame.text
                for shape in slide.shapes:
                    assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=deck.slide_width+2 and shape.top+shape.height<=deck.slide_height+2
            pictures=[(i,shape) for i,slide in enumerate(deck.slides) for shape in slide.shapes if shape.shape_type==13]
            assert len(pictures)==1 and pictures[0][0]==rows.index(grouped)//count
            deck_checks[name]={'slides':len(deck.slides),'pictures':len(pictures),'grouped_image_slide':pictures[0][0]+1}
        saved={'version':3,'title':payload['title'],'metadata':payload['metadata'],'script':script,'plan':plan,'breakdown':rows,
               'approval':approval['approval'],'breakdownHash':approval['plan_hash'],'images':{grouped['id']:approved},
               'continuity':continuity,'visualContext':style}
        saved_id=post('/api/projects',{'project':saved}).json['id']
        restored=client.get('/api/projects/'+saved_id).json['project']
        assert restored['images'][grouped['id']]['approved'] and restored['images'][grouped['id']]['data'].startswith('data:image/jpeg;base64,')
        audit={'mode':'offline rules plus one uploaded validation frame; no live provider call','rows':len(rows),
               'scenes':len(set(r['scene'] for r in rows)),'groups':sum(bool(r.get('coverage_label')) for r in rows),
               'image_shot':grouped['number'],'image_source':approved['source'],'image_approved':approved['approved'],
               'persistence_restored':True,'columns':list(PRODUCTION.values()),'decks':deck_checks,
               'excel':'all 19 columns and all cells matched; shot values are text','csv':'all cells matched','print_rows':len(printed['shots'])}
        (out/'artifact-audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
        print(json.dumps(audit))


if __name__=='__main__':main()
